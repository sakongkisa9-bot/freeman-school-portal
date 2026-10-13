import re

from openpyxl import load_workbook


HEADER_ALIASES = {
    "name": {
        "name",
        "student name",
        "learner name",
        "learner s name",
        "learner full name",
        "name of learner",
        "full name",
    },
    "assessment_number": {
        "assessment number",
        "assessment no",
        "assessment no.",
        "assessmentnumber",
        "knec assessment number",
    },
    "gender": {"gender", "sex"},
    "phone": {"phone", "phone number", "mobile", "telephone"},
    "stream": {"stream", "class stream"},
}


def _normalize_header(value):
    text = str(value or "").strip().lower()
    return re.sub(r"\s+", " ", re.sub(r"[^a-z0-9. ]", " ", text))


def _cell_text(cell):
    value = getattr(cell, "value", cell)
    if value is None:
        return ""
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        if isinstance(value, float) and not value.is_integer():
            return str(value).strip()
        text = str(int(value))
        number_format = getattr(cell, "number_format", "")
        if re.fullmatch(r"0+", number_format):
            return text.zfill(len(number_format))
        return text
    return str(value).strip()


def read_student_rows(file_path):
    """Read student records from an Excel workbook using its header row."""
    workbook = load_workbook(file_path, read_only=True, data_only=True)
    try:
        for worksheet in workbook.worksheets:
            header_row = None
            columns = {}
            for row_number, row in enumerate(worksheet.iter_rows(min_row=1, max_row=20), 1):
                normalized = [_normalize_header(cell.value) for cell in row]
                candidate = {}
                for field, aliases in HEADER_ALIASES.items():
                    for column_index, header in enumerate(normalized):
                        if header in aliases or (
                            field == "assessment_number"
                            and "assessment" in header
                            and ("number" in header or "no" in header.split())
                        ):
                            candidate[field] = column_index
                            break
                if "name" in candidate and "assessment_number" in candidate:
                    header_row = row_number
                    columns = candidate
                    break

            if header_row is None:
                continue

            students = []
            for row in worksheet.iter_rows(min_row=header_row + 1):
                student = {
                    field: _cell_text(row[index]) if index < len(row) else ""
                    for field, index in columns.items()
                }
                if not student.get("name") and not student.get("assessment_number"):
                    continue
                if not student.get("name") or not student.get("assessment_number"):
                    continue
                students.append(student)
            if students:
                return students
        raise ValueError(
            "No student rows found. The workbook must include Name/Learner Name "
            "and Assessment Number columns."
        )
    finally:
        workbook.close()
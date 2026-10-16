import csv
import io
import json
import os
import sqlite3
import tempfile
import unittest
from unittest.mock import patch

import cloud_portal_v2
from werkzeug.security import generate_password_hash


class CBCCloudFlowTests(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.db_path = os.path.join(self.temp_dir.name, "cbc-test.sqlite")
        self.db_path_patch = patch.object(cloud_portal_v2, "DB_PATH", self.db_path)
        self.db_path_patch.start()
        cloud_portal_v2.init_db()
        connection = cloud_portal_v2.get_db()
        school = connection.execute(
            "INSERT INTO schools (school_name, school_code, password_hash, created_at, portal_open) VALUES (?, ?, ?, ?, 1)",
            ("Test School", "cbc-test", generate_password_hash("pw"), "2026-10-14"),
        )
        self.school_id = school.lastrowid
        connection.execute(
            "INSERT INTO teacher_assignments (school_id, class_name, subject, teacher_name, teacher_code) VALUES (?, 'Grade 3', 'MATH', 'Teacher', 'T001')",
            (self.school_id,),
        )
        connection.execute(
            "INSERT INTO students (school_id, grade, adm_no, student_name, gender, assessment_number) VALUES (?, 'Grade 3', 'A001', 'A Learner', 'F', 'K0010203002')",
            (self.school_id,),
        )
        connection.commit()
        connection.close()
        self.client = cloud_portal_v2.app.test_client()
        with self.client.session_transaction() as session:
            session["school_id"] = self.school_id
            session["teacher_code"] = "T001"
            session["role"] = "admin"

    def tearDown(self):
        self.db_path_patch.stop()
        self.temp_dir.cleanup()

    def _save_curriculum(self):
        return self.client.post(
            "/cbc/curriculum/Grade%203",
            data={
                "curriculum_json": '[{"name":"Mathematics","strands":[{"name":"Number","sub_strands":["Fractions"]}]}]'
            },
        )

    def _seed_global(self, grade, area, strands):
        connection = sqlite3.connect(self.db_path)
        connection.execute(
            "INSERT INTO global_curriculum (grade, area, strands_json, updated_at) VALUES (?, ?, ?, 'now') ON CONFLICT(grade, area) DO UPDATE SET strands_json = excluded.strands_json",
            (grade, area, json.dumps(strands)),
        )
        connection.commit()
        connection.close()
        # Opening the gradebook copies the shared strands into the school's tables.
        self.client.get(f"/cbc/gradebook/{grade.replace(' ', '%20')}?area={area}")

    def test_desktop_curriculum_sync_and_selected_sub_strand_ratings(self):
        curriculum = [
            {
                "name": "Mathematics",
                "strands": [
                    {"name": "Number", "sub_strands": ["Fractions", "Decimals"]}
                ],
            }
        ]
        self._seed_global("Grade 3", "Mathematics", curriculum[0]["strands"])

        connection = sqlite3.connect(self.db_path)
        student_id = connection.execute(
            "SELECT id FROM students WHERE adm_no = 'A001'"
        ).fetchone()[0]
        strand_id = connection.execute(
            "SELECT id FROM cbc_strands WHERE name = 'Number'"
        ).fetchone()[0]
        sub_strands = connection.execute(
            "SELECT id, name FROM cbc_sub_strands WHERE strand_id = ? ORDER BY name",
            (strand_id,),
        ).fetchall()
        connection.close()
        self.assertEqual([row[1] for row in sub_strands], ["Decimals", "Fractions"])

        page = self.client.get(
            f"/cbc/gradebook/Grade%203?area=Mathematics&strand_id={strand_id}"
        )
        self.assertEqual(page.status_code, 200)
        self.assertIn(b"Sub-strand", page.data)
        self.assertIn(b"Fractions", page.data)
        self.assertIn(b"Decimals", page.data)

        with self.client.session_transaction() as session:
            session["role"] = "teacher"
        for sub_strand_id, rating in ((sub_strands[0][0], "2"), (sub_strands[1][0], "4")):
            response = self.client.post(
                f"/cbc/gradebook/Grade%203?area=Mathematics&term=Term%201&academic_year=2026&assessment_type=Formative&strand_id={strand_id}&sub_strand_id={sub_strand_id}",
                data={
                    f"rating_{student_id}_{sub_strand_id}": rating,
                    "strand_id": str(strand_id),
                    "sub_strand_id": str(sub_strand_id),
                    "term": "Term 1",
                    "academic_year": "2026",
                    "assessment_type": "Formative",
                },
            )
            self.assertEqual(response.status_code, 302)

        connection = sqlite3.connect(self.db_path)
        saved = connection.execute(
            "SELECT sub_strand_id, rating FROM cbc_rubric_scores ORDER BY sub_strand_id"
        ).fetchall()
        connection.close()
        self.assertEqual(
            dict(saved), {sub_strands[0][0]: "2", sub_strands[1][0]: "4"}
        )

        fetched = self.client.post(
            "/api/cbc/fetch_ratings",
            json={
                "school_code": "cbc-test",
                "password": "pw",
                "grade": "Grade 3",
                "area": "Mathematics",
                "term": "Term 1",
                "academic_year": 2026,
            },
        )
        self.assertEqual(fetched.status_code, 200)
        payload = fetched.get_json()
        self.assertEqual(payload["columns"][-1], "Level")
        self.assertEqual(len(payload["records"]), 2)
        self.assertEqual({row[4] for row in payload["records"]}, {"Decimals", "Fractions"})

    def _publish_math_curriculum(self):
        self._seed_global(
            "Grade 3", "Mathematics",
            [{"name": "Number", "sub_strands": ["Fractions", "Decimals"]}],
        )
        connection = sqlite3.connect(self.db_path)
        ids = [row[0] for row in connection.execute("SELECT id FROM cbc_sub_strands ORDER BY id")]
        connection.close()
        return ids

    def test_autofill_levels_from_current_exam_percentages(self):
        sub_strand_ids = self._publish_math_curriculum()
        connection = sqlite3.connect(self.db_path)
        connection.execute(
            "INSERT INTO students (school_id, grade, adm_no, student_name, gender, assessment_number) VALUES (?, 'Grade 3', 'A002', 'Second Learner', 'M', 'K0010203003')",
            (self.school_id,),
        )
        for adm_no, name, score in (("A001", "A Learner", "80"), ("A002", "Second Learner", "49")):
            connection.execute(
                "INSERT INTO marks (school_id, grade, adm_no, student_name, exam_title, subject_scores_json, updated_at) VALUES (?, 'Grade 3', ?, ?, 'TERM 1 EXAM', ?, '2026-03-01T00:00:00')",
                (self.school_id, adm_no, name, json.dumps({"MATH": {"score": score}})),
            )
        connection.commit()
        connection.close()

        page = self.client.get("/cbc/gradebook/Grade%203?area=Mathematics")
        self.assertIn(b"TERM 1 EXAM", page.data)

        response = self.client.post(
            "/cbc/gradebook/Grade%203?area=Mathematics&term=Term%201&academic_year=2026",
            data={"action": "autofill", "covered_sub_strand_ids": [str(item) for item in sub_strand_ids]},
        )
        self.assertEqual(response.status_code, 302)
        connection = sqlite3.connect(self.db_path)
        saved = connection.execute(
            "SELECT assessment_number, sub_strand_id, assessment_type, rating FROM cbc_rubric_scores"
        ).fetchall()
        connection.close()
        expected = {
            (number, sub_id, "Formative", rating)
            for number, rating in (("K0010203002", "4"), ("K0010203003", "2"))
            for sub_id in sub_strand_ids
        }
        self.assertEqual(set(saved), expected)

    def test_autofill_without_current_exam_changes_nothing(self):
        sub_strand_ids = self._publish_math_curriculum()
        page = self.client.get("/cbc/gradebook/Grade%203?area=Mathematics")
        self.assertIn(b"No current exam", page.data)
        self.client.post(
            "/cbc/gradebook/Grade%203?area=Mathematics",
            data={"action": "autofill", "covered_sub_strand_ids": [str(sub_strand_ids[0])]},
        )
        connection = sqlite3.connect(self.db_path)
        count = connection.execute("SELECT COUNT(*) FROM cbc_rubric_scores").fetchone()[0]
        connection.close()
        self.assertEqual(count, 0)

    def test_autofill_keeps_existing_ratings(self):
        sub_strand_ids = self._publish_math_curriculum()
        connection = sqlite3.connect(self.db_path)
        connection.execute(
            "INSERT INTO marks (school_id, grade, adm_no, student_name, exam_title, subject_scores_json, updated_at) VALUES (?, 'Grade 3', 'A001', 'A Learner', 'TERM 1 EXAM', ?, '2026-03-01T00:00:00')",
            (self.school_id, json.dumps({"MATH": {"score": "85"}})),
        )
        connection.execute(
            "INSERT INTO cbc_rubric_scores (school_id, grade, assessment_number, academic_year, term, assessment_type, sub_strand_id, rating, updated_at) VALUES (?, 'Grade 3', 'K0010203002', 2026, 'Term 1', 'Formative', ?, '1', 'now')",
            (self.school_id, sub_strand_ids[0]),
        )
        connection.commit()
        connection.close()
        self.client.post(
            "/cbc/gradebook/Grade%203?area=Mathematics&term=Term%201&academic_year=2026",
            data={"action": "autofill", "covered_sub_strand_ids": [str(item) for item in sub_strand_ids]},
        )
        connection = sqlite3.connect(self.db_path)
        saved = dict(connection.execute("SELECT sub_strand_id, rating FROM cbc_rubric_scores").fetchall())
        connection.close()
        self.assertEqual(saved, {sub_strand_ids[0]: "1", sub_strand_ids[1]: "4"})

    def _create_exam(self, title="TERM 1 EXAM", **scope):
        payload = {"school_code": "cbc-test", "password": "pw", "title": title}
        payload.update(scope or {"all_classes": True})
        return self.client.post("/api/create_exam", json=payload)

    def test_dashboard_shows_only_assigned_classes_and_exam_button(self):
        page = self.client.get("/dashboard")
        self.assertIn(b"Grade 3", page.data)
        self.assertNotIn(b"Grade 4", page.data)
        self.assertNotIn(b"Percentage Marks", page.data)
        self.assertIn(b"KNEC CBA", page.data)
        self.assertIn(b"Competency Gradebook", page.data)
        self.assertNotIn(b"/exam/select_subject/", page.data)

        self.assertEqual(self._create_exam("MID TERM").status_code, 200)
        self.assertIn(b"MID TERM", self.client.get("/dashboard").data)

        self._create_exam("OTHER CLASS EXAM", all_classes=False, classes=["Grade 4"])
        self.assertNotIn(b"OTHER CLASS EXAM", self.client.get("/dashboard").data)

    def test_exam_marks_save_review_finish_and_lock(self):
        self._create_exam("MID TERM", all_classes=False, classes=["Grade 3"])
        url = "/exam/marks/Grade%203?subject=MATH"
        self.assertIn(b"/exam/marks/Grade%203?subject=MATH", self.client.get("/exam/select_subject/Grade%203").data)
        self.assertEqual(self.client.get(url).status_code, 200)

        too_high = self.client.post(url, data={"out_of": "50", "score_A001": "51", "action": "save"})
        self.assertIn(b"Marks must be between 0 and 50", too_high.data)

        saved = self.client.post(url, data={"out_of": "50", "score_A001": "40", "action": "save"})
        self.assertEqual(saved.status_code, 302)
        review = self.client.get(url + "&review=1")
        self.assertIn(b"Finish", review.data)
        self.assertIn(b'value="40"', review.data)

        connection = sqlite3.connect(self.db_path)
        entry = json.loads(connection.execute("SELECT subject_scores_json FROM marks WHERE adm_no = 'A001'").fetchone()[0])["MATH"]
        connection.close()
        self.assertEqual((entry["score"], entry["raw"], entry["out_of"]), ("80", "40", 50))

        edited = self.client.post(url, data={"out_of": "50", "score_A001": "25", "action": "finish"})
        self.assertEqual(edited.status_code, 302)
        locked = self.client.post(url, data={"out_of": "50", "score_A001": "50", "action": "save"})
        self.assertEqual(locked.status_code, 302)
        self.assertNotIn(b"MID TERM", self.client.get("/dashboard").data)

        connection = sqlite3.connect(self.db_path)
        score = json.loads(connection.execute("SELECT subject_scores_json FROM marks WHERE adm_no = 'A001'").fetchone()[0])["MATH"]["score"]
        connection.close()
        self.assertEqual(score, "50")

    def test_new_exam_replaces_previous_marks(self):
        self._create_exam("FIRST EXAM")
        self.client.post(
            "/exam/marks/Grade%203?subject=MATH",
            data={"out_of": "100", "score_A001": "70", "action": "save"},
        )
        self._create_exam("SECOND EXAM")
        connection = sqlite3.connect(self.db_path)
        marks = connection.execute("SELECT COUNT(*) FROM marks").fetchone()[0]
        titles = [row[0] for row in connection.execute("SELECT title FROM exams")]
        connection.close()
        self.assertEqual(marks, 0)
        self.assertEqual(titles, ["SECOND EXAM"])

    def test_saving_levels_marks_sub_strand_as_covered(self):
        sub_strand_ids = self._publish_math_curriculum()
        url = "/cbc/gradebook/Grade%203?area=Mathematics&term=Term%201&academic_year=2026"
        before = self.client.get(url)
        self.assertIn(b"0 of 1 strands", before.data)
        self.assertIn(b"0 of 2 sub-strands covered", before.data)

        connection = sqlite3.connect(self.db_path)
        student_id = connection.execute("SELECT id FROM students WHERE adm_no = 'A001'").fetchone()[0]
        connection.close()
        self.client.post(
            url + f"&sub_strand_id={sub_strand_ids[0]}",
            data={f"rating_{student_id}_{sub_strand_ids[0]}": "3"},
        )
        partial = self.client.get(url)
        self.assertIn(b"1 of 2 sub-strands covered", partial.data)
        self.assertIn(b"0 of 1 strands", partial.data)

        self.client.post(
            url + f"&sub_strand_id={sub_strand_ids[1]}",
            data={f"rating_{student_id}_{sub_strand_ids[1]}": "2"},
        )
        done = self.client.get(url)
        self.assertIn(b"1 of 1 strands", done.data)
        self.assertIn(b"2 of 2 sub-strands covered", done.data)

    def test_fetch_covered_strands_lists_only_marked_sub_strands(self):
        sub_strand_ids = self._publish_math_curriculum()
        connection = sqlite3.connect(self.db_path)
        student_id = connection.execute("SELECT id FROM students WHERE adm_no = 'A001'").fetchone()[0]
        connection.close()
        self.client.post(
            f"/cbc/gradebook/Grade%203?area=Mathematics&term=Term%202&academic_year=2026&sub_strand_id={sub_strand_ids[0]}",
            data={f"rating_{student_id}_{sub_strand_ids[0]}": "3"},
        )
        request = {"school_code": "cbc-test", "password": "pw", "grade": "Grade 3", "area": "Mathematics", "academic_year": 2026}
        payload = self.client.post("/api/cbc/fetch_covered_strands", json=request).get_json()
        self.assertTrue(payload["success"])
        self.assertEqual(payload["strands"], [{
            "name": "Number",
            "sub_strands": [
                {"name": "Fractions", "terms": ["Term 2"]},
                {"name": "Decimals", "terms": []},
            ],
        }])

        wrong = self.client.post("/api/cbc/fetch_covered_strands", json={**request, "password": "bad"})
        self.assertEqual(wrong.status_code, 401)
        other_subject = self.client.post("/api/cbc/fetch_covered_strands", json={**request, "area": "Science"})
        self.assertFalse(other_subject.get_json()["success"])
        next_year = self.client.post("/api/cbc/fetch_covered_strands", json={**request, "academic_year": 2025})
        self.assertEqual(next_year.get_json()["strands"][0]["sub_strands"][0]["terms"], [])

    def test_consume_marks_endpoint_keeps_marks(self):
        self._create_exam("MID TERM")
        self.client.post(
            "/exam/marks/Grade%203?subject=MATH",
            data={"out_of": "100", "score_A001": "70", "action": "save"},
        )
        response = self.client.post(
            "/api/consume_marks",
            json={"school_code": "cbc-test", "grade": "Grade 3"},
        )
        self.assertTrue(response.get_json()["success"])
        connection = sqlite3.connect(self.db_path)
        count = connection.execute("SELECT COUNT(*) FROM marks").fetchone()[0]
        connection.close()
        self.assertEqual(count, 1)

    def _add_second_learner(self):
        connection = sqlite3.connect(self.db_path)
        connection.execute(
            "INSERT INTO students (school_id, grade, adm_no, student_name, gender, assessment_number) VALUES (?, 'Grade 3', 'A002', 'Second Learner', 'M', 'K0010203003')",
            (self.school_id,),
        )
        connection.commit()
        connection.close()

    def test_blank_exam_marks_become_ab_and_zero_needs_confirmation(self):
        self._add_second_learner()
        self._create_exam("MID TERM", all_classes=False, classes=["Grade 3"])
        url = "/exam/marks/Grade%203?subject=MATH"
        base = {"out_of": "50", "action": "save"}

        unconfirmed = self.client.post(url, data={**base, "score_A001": "0", "score_A002": "30"})
        self.assertIn(b"Confirm the zero marks for A Learner", unconfirmed.data)
        connection = sqlite3.connect(self.db_path)
        self.assertEqual(connection.execute("SELECT COUNT(*) FROM marks").fetchone()[0], 0)
        connection.close()

        self.client.post(url, data={**base, "score_A001": "0", "score_A002": "30", "confirm_zero": "1"})
        self.client.post(url, data={**base, "score_A001": "40", "score_A002": ""})
        connection = sqlite3.connect(self.db_path)
        scores = {
            row[0]: json.loads(row[1])["MATH"]
            for row in connection.execute("SELECT adm_no, subject_scores_json FROM marks")
        }
        connection.close()
        self.assertEqual(scores["A001"]["score"], "80")
        self.assertEqual((scores["A002"]["score"], scores["A002"]["raw"]), ("AB", "AB"))
        review = self.client.get(url + "&review=1")
        self.assertIn(b'value="AB"', review.data)

    def test_autofill_writes_ab_for_absent_learners_and_warns(self):
        sub_strand_ids = self._publish_math_curriculum()
        self._add_second_learner()
        connection = sqlite3.connect(self.db_path)
        for adm_no, name, score in (("A001", "A Learner", "85"), ("A002", "Second Learner", "AB")):
            connection.execute(
                "INSERT INTO marks (school_id, grade, adm_no, student_name, exam_title, subject_scores_json, updated_at) VALUES (?, 'Grade 3', ?, ?, 'TERM 1 EXAM', ?, '2026-03-01T00:00:00')",
                (self.school_id, adm_no, name, json.dumps({"MATH": {"score": score}})),
            )
        connection.commit()
        connection.close()

        page = self.client.get("/cbc/gradebook/Grade%203?area=Mathematics")
        self.assertIn(b"Absent (AB) in this exam", page.data)
        self.assertIn(b"Second Learner", page.data)
        self.assertIn(b'<option value="AB">AB</option>', page.data)

        response = self.client.post(
            "/cbc/gradebook/Grade%203?area=Mathematics&term=Term%201&academic_year=2026",
            data={"action": "autofill", "covered_sub_strand_ids": [str(item) for item in sub_strand_ids]},
            follow_redirects=True,
        )
        self.assertIn(b"Absent (AB) was written for all selected sub-strands of: Second Learner", response.data)
        connection = sqlite3.connect(self.db_path)
        saved = {
            (number, rating)
            for number, rating in connection.execute(
                "SELECT assessment_number, rating FROM cbc_rubric_scores"
            )
        }
        connection.close()
        self.assertEqual(saved, {("K0010203002", "4"), ("K0010203003", "AB")})

    def test_developer_page_is_password_protected_and_stores_shared_strands(self):
        strands = json.dumps([{"name": "1.0 Conservation", "sub_strands": ["1.1 Soil"]}])
        form = {"grade": "Grade 4", "area": "Agriculture", "strands_json": strands}
        with patch.dict(os.environ, {}):
            key = cloud_portal_v2.SYSTEM_ADMIN_KEY
            locked = self.client.get("/developer/strands")
            self.assertIn(b"Developer password", locked.data)
            self.client.post("/developer/strands/save", data=form)
            self.client.post("/developer/strands", data={"action": "login", "password": "wrong"})
            self.assertIn(b"Developer password", self.client.get("/developer/strands").data)

            connection = sqlite3.connect(self.db_path)
            self.assertEqual(connection.execute("SELECT COUNT(*) FROM global_curriculum").fetchone()[0], 0)
            connection.close()

            self.client.post("/developer/strands", data={"action": "login", "password": key})
            self.client.post("/developer/strands/save", data=form)
            page = self.client.get("/developer/strands")
            self.assertIn(b"Agriculture", page.data)

            connection = sqlite3.connect(self.db_path)
            row = connection.execute("SELECT id, strands_json FROM global_curriculum").fetchone()
            connection.close()
            self.assertEqual(json.loads(row[1])[0]["sub_strands"], ["1.1 Soil"])

            self.client.post("/developer/strands/delete", data={"id": str(row[0])})
            connection = sqlite3.connect(self.db_path)
            self.assertEqual(connection.execute("SELECT COUNT(*) FROM global_curriculum").fetchone()[0], 0)
            connection.close()

    def test_developer_upload_reads_strands_from_curriculum_pdf(self):
        from reportlab.lib.pagesizes import A4
        from reportlab.platypus import SimpleDocTemplate, Table, TableStyle
        from reportlab.lib import colors

        buffer = io.BytesIO()
        table = Table([
            ["Strands", "Sub Strands", "Suggested Number of Lessons"],
            ["1.0 Conservation of Resources", "1.1 Soil Conservation", "7"],
            ["", "1.2 Water Conservation", "7"],
            ["2.0 Food Production", "2.1 Growing Fruits", "16"],
        ])
        table.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.black)]))
        SimpleDocTemplate(buffer, pagesize=A4).build([table])
        buffer.seek(0)

        with patch.dict(os.environ, {}):
            self.client.post("/developer/strands", data={"action": "login", "password": cloud_portal_v2.SYSTEM_ADMIN_KEY})
            response = self.client.post(
                "/developer/strands/preview",
                data={"grade": "Grade 4", "area": "Agriculture", "pdf": (buffer, "agri.pdf")},
                content_type="multipart/form-data",
            )
        self.assertEqual(response.status_code, 200)
        for text in (b"1.0 Conservation of Resources", b"1.2 Water Conservation", b"2.1 Growing Fruits", b"Save to Portal"):
            self.assertIn(text, response.data)

    def test_developer_unlock_is_checked_on_the_server_and_key_is_not_in_pages(self):
        key = cloud_portal_v2.SYSTEM_ADMIN_KEY
        for url in ("/", "/register"):
            self.assertNotIn(key.encode(), self.client.get(url).data)
        for url in ("/register", "/admin/schools", "/admin/schools/x/edit"):
            self.assertEqual(self.client.get(url).status_code, 302)
        self.assertEqual(self.client.post("/admin/schools/x/delete").status_code, 302)
        self.assertEqual(self.client.post("/register", data={"master_key": key}).status_code, 302)
        wrong = self.client.post("/developer/unlock", json={"password": "nope"})
        self.assertEqual(wrong.status_code, 401)
        right = self.client.post("/developer/unlock", json={"password": key})
        self.assertTrue(right.get_json()["success"])
        unlocked = self.client.get("/").data
        self.assertIn(b'<div id="dev-options" style="display: block', unlocked)
        self.assertIn(b'type="password" id="dev-password"', unlocked)
        self.assertEqual(self.client.get("/register").status_code, 200)
        self.assertEqual(self.client.get("/admin/schools").status_code, 200)
        self.assertNotIn(key.encode(), self.client.get("/register").data)

    def test_strands_api_requires_shared_curriculum_and_valid_login(self):
        request = {"school_code": "cbc-test", "password": "pw", "grade": "Grade 3", "area": "Mathematics", "academic_year": 2026}
        self.assertFalse(self.client.post("/api/cbc/fetch_covered_strands", json=request).get_json()["success"])
        bad = self.client.post("/api/cbc/fetch_covered_strands", json={**request, "password": "wrong"})
        self.assertEqual(bad.status_code, 401)

    def test_curriculum_gradebook_and_csv_staging_use_assessment_number(self):
        curriculum_response = self._save_curriculum()
        self.assertEqual(curriculum_response.status_code, 200)
        with self.client.session_transaction() as session:
            session["role"] = "teacher"

        page = self.client.get(
            "/cbc/gradebook/Grade%203?area=Mathematics&term=Term%201&academic_year=2026"
        )
        self.assertEqual(page.status_code, 200)
        self.assertIn(b"K0010203002", page.data)

        connection = sqlite3.connect(self.db_path)
        student_id = connection.execute("SELECT id FROM students WHERE adm_no='A001'").fetchone()[0]
        sub_strand_id = connection.execute("SELECT id FROM cbc_sub_strands WHERE name='Fractions'").fetchone()[0]
        connection.close()
        for value in ("4",):
            response = self.client.post(
                "/cbc/gradebook/Grade%203?area=Mathematics&term=Term%201&academic_year=2026",
                data={"rating_{}_{}".format(student_id, sub_strand_id): value},
            )
            self.assertEqual(response.status_code, 302)

        progress = self.client.get(
            "/cbc/gradebook/Grade%203?area=Mathematics&term=Term%201&academic_year=2026"
        )
        self.assertNotIn(b"Localized Summative", progress.data)
        self.assertIn(b"EE", progress.data)
        profile = self.client.get(
            "/cbc/profile/Grade%203/K0010203002?term=Term%201&academic_year=2026"
        )
        self.assertEqual(profile.status_code, 200)
        self.assertIn(b"Fractions", profile.data)
        self.assertIn(b"EE1 / EE2", profile.data)
        self.assertNotIn(b"Total Score", profile.data)
        self.assertNotIn(b"Rank", profile.data)
        connection = sqlite3.connect(self.db_path)
        ratings = connection.execute(
            "SELECT assessment_number, assessment_type, rating FROM cbc_rubric_scores ORDER BY assessment_type"
        ).fetchall()
        connection.execute(
            "INSERT INTO cba_term_task_configs (school_id, grade, subject, term, tasks_json, updated_at) VALUES (?, 'Grade 3', 'MATH', 'Term 1', ?, 'now')",
            (
                self.school_id,
                '[{"name":"Task 1","ranges":{"Level 1":{"min":0,"max":2},"Level 2":{"min":3,"max":5},"Level 3":{"min":6,"max":8},"Level 4":{"min":9,"max":10}}},{"name":"Task 2","ranges":{"Level 1":{"min":0,"max":2},"Level 2":{"min":3,"max":5},"Level 3":{"min":6,"max":8},"Level 4":{"min":9,"max":10}}}]',
            ),
        )
        connection.commit()
        connection.close()
        self.assertEqual(ratings, [("K0010203002", "Formative", "4")])

        staging = self.client.get("/cba/staging/Grade%203?subject=MATH&term=Term%201")
        self.assertEqual(staging.status_code, 200)
        download = self.client.get("/cba/staging/Grade%203/download?subject=MATH&term=Term%201")
        self.assertEqual(download.status_code, 200)
        rows = list(csv.reader(io.StringIO(download.data.decode("utf-8-sig"))))
        self.assertEqual(rows[0], ["Assessment Number", "Learner's Name", "Gender", "Task 1", "Task 2"])
        self.assertEqual(rows[1][:3], ["K0010203002", "A Learner", "F"])

    def test_csv_staging_blocks_missing_assessment_numbers(self):
        self._save_curriculum()
        connection = sqlite3.connect(self.db_path)
        connection.execute(
            "INSERT INTO students (school_id, grade, adm_no, student_name, gender, assessment_number) VALUES (?, 'Grade 3', 'A002', 'Missing Number', 'M', '')",
            (self.school_id,),
        )
        connection.execute(
            "INSERT INTO cba_term_task_configs (school_id, grade, subject, term, tasks_json, updated_at) VALUES (?, 'Grade 3', 'MATH', 'Term 1', ?, 'now')",
            (self.school_id, '[{"name":"Task 1","ranges":{"Level 1":{"min":0,"max":2},"Level 2":{"min":3,"max":5},"Level 3":{"min":6,"max":8},"Level 4":{"min":9,"max":10}}}]'),
        )
        connection.commit()
        connection.close()
        with self.client.session_transaction() as session:
            session["role"] = "teacher"
        response = self.client.get("/cba/staging/Grade%203/download?subject=MATH&term=Term%201")
        self.assertEqual(response.status_code, 400)
        self.assertIn("Missing Number", response.get_json()["missing_assessment_numbers"])

    def test_sync_api_uses_assessment_number_without_aggregate_total(self):
        tasks = [{
            "name": "Task 1",
            "ranges": {
                "Level 1": {"min": 0, "max": 2},
                "Level 2": {"min": 3, "max": 5},
                "Level 3": {"min": 6, "max": 8},
                "Level 4": {"min": 9, "max": 10},
            },
        }]
        publish = self.client.post(
            "/api/upload_cba_tasks",
            json={
                "school_code": "cbc-test",
                "username": "admin",
                "password": "pw",
                "grade": "Grade 3",
                "subject": "MATH",
                "term": "Term 1",
                "tasks": tasks,
            },
        )
        self.assertEqual(publish.status_code, 200)
        connection = sqlite3.connect(self.db_path)
        connection.execute(
            "INSERT INTO cba_term_scores (school_id, grade, subject, term, adm_no, student_name, task_name, raw_score, teacher_code, updated_at) VALUES (?, 'Grade 3', 'MATH', 'Term 1', 'A001', 'A Learner', 'Task 1', '4', 'T001', 'now')",
            (self.school_id,),
        )
        connection.commit()
        connection.close()
        response = self.client.post(
            "/api/fetch_cba_results",
            json={
                "school_code": "cbc-test",
                "password": "pw",
                "grade": "Grade 3",
                "subject": "MATH",
                "term": "Term 1",
            },
        )
        self.assertEqual(response.status_code, 200)
        payload = response.get_json()
        self.assertEqual(payload["columns"], ["Assessment Number", "Learner's Name", "Gender", "Task 1"])
        self.assertEqual(payload["records"][0], ["K0010203002", "A Learner", "F", 2])

    def test_mixed_level_and_max_tasks_export_levels_and_raw_scores(self):
        tasks = [
            {
                "name": "Task 1",
                "ranges": {
                    "Level 1": {"min": 0, "max": 2},
                    "Level 2": {"min": 3, "max": 5},
                    "Level 3": {"min": 6, "max": 8},
                    "Level 4": {"min": 9, "max": 10},
                },
            },
            {"name": "Task 2", "mode": "max", "max_score": 8},
        ]
        publish = self.client.post(
            "/api/upload_cba_tasks",
            json={
                "school_code": "cbc-test",
                "username": "admin",
                "password": "pw",
                "grade": "Grade 3",
                "subject": "MATH",
                "term": "Term 1",
                "tasks": tasks,
            },
        )
        self.assertEqual(publish.status_code, 200)
        with self.client.session_transaction() as session:
            session["role"] = "teacher"
        page = self.client.get("/cba/marks/Grade%203?subject=MATH&term=Term%201")
        self.assertIn(b'id="taskSelect"', page.data)
        self.assertIn(b"max 8", page.data)

        student_id = sqlite3.connect(self.db_path).execute(
            "SELECT id FROM students WHERE adm_no='A001'"
        ).fetchone()[0]
        too_high = self.client.post(
            "/cba/marks/Grade%203?subject=MATH&term=Term%201",
            data={f"score_{student_id}_0": "4", f"score_{student_id}_1": "9"},
        )
        self.assertEqual(too_high.status_code, 302)
        connection = sqlite3.connect(self.db_path)
        self.assertEqual(connection.execute("SELECT COUNT(*) FROM cba_term_scores").fetchone()[0], 0)
        connection.close()

        self.client.post(
            "/cba/marks/Grade%203?subject=MATH&term=Term%201",
            data={f"score_{student_id}_0": "4", f"score_{student_id}_1": "7"},
        )
        csv_response = self.client.get(
            "/cba/staging/Grade%203/download?subject=MATH&term=Term%201"
        )
        rows = list(csv.reader(io.StringIO(csv_response.data.decode("utf-8-sig"))))
        self.assertEqual(rows[1], ["K0010203002", "A Learner", "F", "2", "7"])

        excel = self.client.get("/cba/export/Grade%203/xlsx?subject=MATH&term=Term%201")
        from openpyxl import load_workbook

        sheet = load_workbook(io.BytesIO(excel.data), data_only=True).active
        self.assertEqual([sheet.cell(3, c).value for c in (4, 5)], [2, 7])

        fetched = self.client.post(
            "/api/fetch_cba_results",
            json={"school_code": "cbc-test", "password": "pw", "grade": "Grade 3", "subject": "MATH", "term": "Term 1"},
        )
        self.assertEqual(fetched.get_json()["records"][0][3:], [2, 7])

    def test_milestone_route_applies_requested_weighting(self):
        connection = sqlite3.connect(self.db_path)
        connection.execute(
            "INSERT INTO students (school_id, grade, adm_no, student_name, gender, assessment_number) VALUES (?, 'Grade 9', 'A009', 'A Learner', 'F', 'K0010203002')",
            (self.school_id,),
        )
        connection.commit()
        connection.close()
        response = self.client.post(
            "/api/cbc/milestone-results",
            json={
                "grade": "Grade 9",
                "assessment_number": "K0010203002",
                "academic_year": 2026,
                "grade7_sba_percent": 80,
                "grade8_sba_percent": 90,
                "national_percent": 90,
            },
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.get_json()["weighted_result"], 88.0)

    def test_start_new_year_clears_portal_exam_and_advances_period(self):
        connection = sqlite3.connect(self.db_path)
        connection.execute("UPDATE schools SET current_year = 2026 WHERE id = ?", (self.school_id,))
        exam = connection.execute(
            "INSERT INTO exams (school_id, title, scope_json, created_at) VALUES (?, 'Term Exam', '{}', 'now')",
            (self.school_id,),
        )
        connection.execute(
            "INSERT INTO exam_submissions (school_id, exam_id, grade, subject, out_of, status, updated_at) VALUES (?, ?, 'Grade 3', 'MATH', 100, 'final', 'now')",
            (self.school_id, exam.lastrowid),
        )
        connection.execute(
            "INSERT INTO marks (school_id, grade, adm_no, student_name, exam_title, subject_scores_json, total_points, average_level, updated_at) VALUES (?, 'Grade 3', 'A001', 'A Learner', 'Term Exam', '{}', '0', '', 'now')",
            (self.school_id,),
        )
        connection.commit()
        connection.close()

        response = self.client.post(
            "/api/start_new_year",
            json={"school_code": "cbc-test", "password": "pw", "academic_year": 2027, "term": "Term 1"},
        )
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.get_json()["success"])

        connection = sqlite3.connect(self.db_path)
        period = connection.execute(
            "SELECT current_year, current_term FROM schools WHERE id = ?", (self.school_id,)
        ).fetchone()
        remaining = connection.execute(
            "SELECT (SELECT COUNT(*) FROM marks WHERE school_id = ?), (SELECT COUNT(*) FROM exams WHERE school_id = ?), (SELECT COUNT(*) FROM exam_submissions WHERE school_id = ?)",
            (self.school_id, self.school_id, self.school_id),
        ).fetchone()
        connection.close()
        self.assertEqual(period, (2027, "Term 1"))
        self.assertEqual(remaining, (0, 0, 0))

        backwards = self.client.post(
            "/api/start_new_year",
            json={"school_code": "cbc-test", "password": "pw", "academic_year": 2026, "term": "Term 1"},
        )
        self.assertEqual(backwards.status_code, 409)


class CBCRubricSchemaUpgradeTests(unittest.TestCase):
    def test_legacy_ratings_are_preserved_as_formative(self):
        temp_dir = tempfile.TemporaryDirectory()
        db_path = os.path.join(temp_dir.name, "legacy-cbc.sqlite")
        patcher = patch.object(cloud_portal_v2, "DB_PATH", db_path)
        patcher.start()
        try:
            connection = sqlite3.connect(db_path)
            connection.executescript(
                """
                CREATE TABLE cbc_learning_areas (id INTEGER PRIMARY KEY, school_id INTEGER, grade TEXT, name TEXT);
                CREATE TABLE cbc_strands (id INTEGER PRIMARY KEY, learning_area_id INTEGER, name TEXT);
                CREATE TABLE cbc_sub_strands (id INTEGER PRIMARY KEY, strand_id INTEGER, name TEXT);
                CREATE TABLE cbc_rubric_scores (
                    id INTEGER PRIMARY KEY, school_id INTEGER, grade TEXT, assessment_number TEXT,
                    academic_year INTEGER, term TEXT, sub_strand_id INTEGER, rating TEXT,
                    evidence TEXT, teacher_code TEXT, updated_at TEXT
                );
                INSERT INTO cbc_learning_areas VALUES (1, 1, 'Grade 3', 'Mathematics');
                INSERT INTO cbc_strands VALUES (1, 1, 'Number');
                INSERT INTO cbc_sub_strands VALUES (1, 1, 'Fractions');
                INSERT INTO cbc_rubric_scores VALUES (1, 1, 'Grade 3', 'K0010203002', 2026, 'Term 1', 1, '4', NULL, 'T001', 'now');
                """
            )
            connection.close()
            cloud_portal_v2.init_db()
            connection = sqlite3.connect(db_path)
            upgraded = connection.execute(
                "SELECT assessment_number, assessment_type, rating FROM cbc_rubric_scores"
            ).fetchone()
            connection.close()
            self.assertEqual(upgraded, ("K0010203002", "Formative", "4"))
        finally:
            patcher.stop()
            temp_dir.cleanup()


if __name__ == "__main__":
    unittest.main()
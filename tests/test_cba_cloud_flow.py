import io
import json
import os
import sqlite3
import tempfile
import unittest
from unittest.mock import patch

import cloud_portal_v2
from openpyxl import load_workbook
from werkzeug.security import generate_password_hash


class CBACloudFlowTests(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.db_path = os.path.join(self.temp_dir.name, "cba-test.sqlite")
        self.db_path_patch = patch.object(cloud_portal_v2, "DB_PATH", self.db_path)
        self.db_path_patch.start()
        cloud_portal_v2.init_db()

        connection = cloud_portal_v2.get_db()
        cursor = connection.execute(
            """
            INSERT INTO schools (school_name, school_code, password_hash, created_at, portal_open)
            VALUES (?, ?, ?, ?, 1)
            """,
            (
                "Test School",
                "test-school",
                generate_password_hash("school-password"),
                "2026-10-13",
            ),
        )
        self.school_id = cursor.lastrowid
        connection.execute(
            """
            INSERT INTO teacher_assignments
                (school_id, class_name, subject, teacher_name, teacher_code)
            VALUES (?, 'Grade 3', 'ENG', 'Teacher One', 'T001')
            """,
            (self.school_id,),
        )
        connection.execute(
            """
            INSERT INTO students (school_id, grade, adm_no, student_name, gender, assessment_number)
            VALUES (?, 'Grade 3', 'A001', 'Amina Wanjiku', 'F', 'K0010203001')
            """,
            (self.school_id,),
        )
        connection.execute(
            """
            INSERT INTO students (school_id, grade, adm_no, student_name, gender, assessment_number)
            VALUES (?, 'Grade 3', 'A002', 'Brian Otieno', 'M', 'K0010203002')
            """,
            (self.school_id,),
        )
        connection.execute(
            """
            INSERT INTO students (school_id, grade, adm_no, student_name, gender, assessment_number)
            VALUES (?, 'Grade 3', 'A003', 'Carol Njeri', 'F', 'K0010203003')
            """,
            (self.school_id,),
        )
        connection.commit()
        connection.close()
        self.client = cloud_portal_v2.app.test_client()
        with self.client.session_transaction() as session:
            session["school_id"] = self.school_id
            session["teacher_code"] = "T001"

    def tearDown(self):
        self.db_path_patch.stop()
        self.temp_dir.cleanup()

    def test_publish_save_convert_and_download_cba_results(self):
        tasks = [
            {
                "name": "Task 1",
                "unit": "words",
                "ranges": {
                    "Level 1": {"type": "range", "min": 0, "max": 44},
                    "Level 2": {"type": "range", "min": 45, "max": 89},
                    "Level 3": {"type": "exact", "value": 90},
                    "Level 4": {"type": "above", "value": 90},
                },
            }
        ]
        publish_response = self.client.post(
            "/api/upload_cba_tasks",
            json={
                "school_code": "test-school",
                "username": "admin",
                "password": "school-password",
                "grade": "Grade 3",
                "subject": "ENG",
                "term": "Term 3",
                "tasks": tasks,
            },
        )
        self.assertEqual(publish_response.status_code, 200)

        subject_response = self.client.get("/cba/select_subject/Grade%203")
        self.assertEqual(subject_response.status_code, 200)
        self.assertIn(b"ENG", subject_response.data)
        self.assertIn(b"/cba/marks/Grade%203?subject=ENG", subject_response.data)

        page_response = self.client.get("/cba/marks/Grade%203?subject=ENG&term=Term%203")
        self.assertEqual(page_response.status_code, 200)
        self.assertIn(b'<th class="student-name">Learner</th>', page_response.data)
        self.assertNotIn(b"Admission No.", page_response.data)
        self.assertIn(b"Term 3", page_response.data)
        self.assertIn(b"Task 1", page_response.data)
        self.assertIn(b"Min 0", page_response.data)
        self.assertIn(b"No limit words", page_response.data)
        self.assertIn(b"studentSearch", page_response.data)

        save_response = self.client.post(
            "/cba/marks/Grade%203?subject=ENG&term=Term%203",
            # A missing field for student A003 must be saved as absent (AB).
            data={"score_1_0": "4", "score_2_0": "90"},
            follow_redirects=True,
        )
        self.assertEqual(save_response.status_code, 200)

        connection = sqlite3.connect(self.db_path)
        raw_score = connection.execute(
            "SELECT raw_score FROM cba_term_scores WHERE adm_no = 'A001' AND task_name = 'Task 1' AND term = 'Term 3'"
        ).fetchone()[0]
        absent_score = connection.execute(
            "SELECT raw_score FROM cba_term_scores WHERE adm_no = 'A003' AND task_name = 'Task 1' AND term = 'Term 3'"
        ).fetchone()[0]
        percentage_count = connection.execute("SELECT COUNT(*) FROM marks").fetchone()[0]
        connection.close()
        self.assertEqual(raw_score, "4")
        self.assertEqual(absent_score, "AB")
        self.assertEqual(percentage_count, 0)

        fetched_results = self.client.post(
            "/api/fetch_cba_results",
            json={
                "school_code": "test-school",
                "username": "admin",
                "password": "school-password",
                "grade": "Grade 3",
                "subject": "ENG",
                "term": "Term 3",
            },
        )
        self.assertEqual(fetched_results.status_code, 200)
        fetched_data = fetched_results.get_json()
        self.assertEqual(
            [row[1] for row in fetched_data["records"]],
            ["Amina Wanjiku", "Brian Otieno", "Carol Njeri"],
        )
        self.assertEqual(fetched_data["records"][2][-1:], ["A"])
        self.assertNotIn("Total Score", fetched_data["columns"])

        new_term_response = self.client.post(
            "/api/upload_cba_tasks",
            json={
                "school_code": "test-school",
                "username": "admin",
                "password": "school-password",
                "grade": "Grade 3",
                "subject": "ENG",
                "term": "Term 2",
                "tasks": tasks,
            },
        )
        self.assertEqual(new_term_response.status_code, 200)
        latest_term_page = self.client.get("/cba/marks/Grade%203?subject=ENG")
        self.assertIn(b"Term 2", latest_term_page.data)
        old_term_page = self.client.get(
            "/cba/marks/Grade%203?subject=ENG&term=Term%203"
        )
        self.assertIn(b'value="4"', old_term_page.data)
        connection = sqlite3.connect(self.db_path)
        old_term_scores = connection.execute(
            "SELECT COUNT(*) FROM cba_term_scores WHERE term = 'Term 3'"
        ).fetchone()[0]
        connection.close()
        self.assertEqual(old_term_scores, 3)

        invalid_response = self.client.post(
            "/cba/marks/Grade%203?subject=ENG&term=Term%203",
            data={"score_1_0": "-1", "score_2_0": "90"},
            follow_redirects=True,
        )
        self.assertEqual(invalid_response.status_code, 200)
        connection = sqlite3.connect(self.db_path)
        preserved_score = connection.execute(
            "SELECT raw_score FROM cba_term_scores WHERE adm_no = 'A001' AND task_name = 'Task 1' AND term = 'Term 3'"
        ).fetchone()[0]
        connection.close()
        self.assertEqual(preserved_score, "4")

        excel_response = self.client.get(
            "/cba/export/Grade%203/xlsx?subject=ENG&term=Term%203"
        )
        self.assertEqual(excel_response.status_code, 200)
        from openpyxl import load_workbook

        workbook = load_workbook(io.BytesIO(excel_response.data), data_only=True)
        worksheet = workbook.active
        self.assertIn("Term 3", worksheet.cell(1, 1).value)
        self.assertEqual(worksheet.cell(2, 1).value, "Assessment Number")
        self.assertEqual(worksheet.cell(2, 2).value, "Learner's Name")
        self.assertEqual(worksheet.cell(2, 3).value, "Gender")
        self.assertEqual(worksheet.cell(2, 4).value, "Task 1")
        self.assertEqual(worksheet.max_column, 4)
        self.assertEqual(worksheet.cell(3, 1).value, "K0010203001")
        self.assertEqual(worksheet.cell(3, 3).value, "F")
        self.assertEqual(worksheet.cell(3, 4).value, 1)
        self.assertEqual(worksheet.cell(4, 4).value, 3)
        self.assertEqual(worksheet.cell(5, 4).value, "A")

        excel_with_legacy_total_flag = self.client.get(
            "/cba/export/Grade%203/xlsx?subject=ENG&term=Term%203&include_total=1"
        )
        self.assertEqual(excel_with_legacy_total_flag.status_code, 200)
        without_total_sheet = load_workbook(
            io.BytesIO(excel_with_legacy_total_flag.data), data_only=True
        ).active
        self.assertEqual(without_total_sheet.cell(2, 4).value, "Task 1")
        self.assertEqual(without_total_sheet.max_column, 4)

        pdf_response = self.client.get(
            "/cba/export/Grade%203/pdf?subject=ENG&term=Term%203"
        )
        self.assertEqual(pdf_response.status_code, 200)
        self.assertTrue(pdf_response.data.startswith(b"%PDF"))

    def test_excel_export_contains_no_aggregate_total(self):
        tasks = [{
            "name": "Task 1",
            "ranges": {
                "Level 1": {"min": 0, "max": 2},
                "Level 2": {"min": 3, "max": 5},
                "Level 3": {"min": 6, "max": 8},
                "Level 4": {"min": 9, "max": 10},
            },
        }]
        connection = sqlite3.connect(self.db_path)
        connection.execute(
            "INSERT INTO cba_term_task_configs (school_id, grade, subject, term, tasks_json, updated_at) VALUES (?, 'Grade 3', 'ENG', 'Term 1', ?, 'now')",
            (self.school_id, json.dumps(tasks)),
        )
        connection.execute(
            "INSERT INTO cba_term_scores (school_id, grade, subject, term, adm_no, student_name, task_name, raw_score, teacher_code, updated_at) VALUES (?, 'Grade 3', 'ENG', 'Term 1', 'A001', 'Amina Wanjiku', 'Task 1', '4', 'T001', 'now')",
            (self.school_id,),
        )
        connection.commit()
        connection.close()
        response = self.client.get("/cba/export/Grade%203/xlsx?subject=ENG&term=Term%201")
        self.assertEqual(response.status_code, 200)
        worksheet = load_workbook(io.BytesIO(response.data), data_only=True).active
        self.assertEqual(
            [worksheet.cell(2, column).value for column in range(1, worksheet.max_column + 1)],
            ["Assessment Number", "Learner's Name", "Gender", "Task 1"],
        )


if __name__ == "__main__":
    unittest.main()
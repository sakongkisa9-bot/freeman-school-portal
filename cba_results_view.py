import re
from tkinter import filedialog, messagebox, ttk

import customtkinter as ctk


class CBAResultsView(ctk.CTkFrame):
    def __init__(self, parent, result, return_callback=None):
        super().__init__(parent)
        self.grade = result["grade"]
        self.subject = result["subject"]
        self.term = result["term"]
        self.columns = list(result["columns"])
        self.records = [list(row) for row in result["records"]]
        self.has_total_score = bool(
            self.columns and self.columns[-1].casefold() == "total score"
        )
        self.report_title = result.get("report_title", "Converted CBA Results")
        self.filename_suffix = result.get("filename_suffix", "converted_cba")
        self.return_callback = return_callback
        self._build_widgets()

    def _build_widgets(self):
        self.grid_columnconfigure(0, weight=1)
        self.grid_rowconfigure(1, weight=1)

        header = ctk.CTkFrame(self, fg_color="transparent")
        header.grid(row=0, column=0, sticky="ew", padx=18, pady=14)
        header.grid_columnconfigure(0, weight=1)
        ctk.CTkLabel(
            header,
            text=f"{self.report_title}  |  {self.grade} / {self.subject} / {self.term}",
            font=("Arial Bold", 20),
            anchor="w",
        ).grid(row=0, column=0, sticky="w")
        ctk.CTkButton(
            header,
            text="Export Excel",
            command=self.export_excel,
        ).grid(row=0, column=1, padx=4)
        ctk.CTkButton(
            header,
            text="Export PDF",
            command=self.export_pdf,
        ).grid(row=0, column=2, padx=4)
        if self.return_callback:
            ctk.CTkButton(
                header, text="Back", command=self.return_callback
            ).grid(row=0, column=3, padx=(4, 0))

        table_frame = ctk.CTkFrame(self)
        table_frame.grid(row=1, column=0, sticky="nsew", padx=18, pady=(0, 18))
        table_frame.grid_columnconfigure(0, weight=1)
        table_frame.grid_rowconfigure(0, weight=1)

        style = ttk.Style()
        style.configure("CBA.Treeview", rowheight=30, font=("Segoe UI", 10))
        style.configure(
            "CBA.Treeview.Heading",
            font=("Segoe UI", 10, "bold"),
            background="#245b4b",
            foreground="white",
        )
        self.tree = ttk.Treeview(
            table_frame,
            columns=[f"column_{index}" for index in range(len(self.columns))],
            show="headings",
            style="CBA.Treeview",
        )
        vertical = ttk.Scrollbar(table_frame, orient="vertical", command=self.tree.yview)
        horizontal = ttk.Scrollbar(table_frame, orient="horizontal", command=self.tree.xview)
        self.tree.configure(yscrollcommand=vertical.set, xscrollcommand=horizontal.set)
        self.tree.grid(row=0, column=0, sticky="nsew")
        vertical.grid(row=0, column=1, sticky="ns")
        horizontal.grid(row=1, column=0, sticky="ew")

        for index, column in enumerate(self.columns):
            identifier = f"column_{index}"
            self.tree.heading(identifier, text=column)
            self.tree.column(
                identifier,
                width=190 if index < 2 else 110,
                minwidth=90,
                stretch=index < 2,
                anchor="w" if index < 2 else "center",
            )
        for record in self.records:
            self.tree.insert("", "end", values=record)

    def _ask_export_path(self, extension):
        include_total = True
        if self.has_total_score:
            include_total = messagebox.askyesno(
                "Include Total Score",
                "Include the Total Score column in this printout?",
                parent=self,
            )
        filename = self._safe_filename(
            f"{self.grade}_{self.subject}_{self.term}_{self.filename_suffix}.{extension}"
        )
        path = filedialog.asksaveasfilename(
            parent=self,
            title=f"Save {self.report_title} {extension.upper()}",
            defaultextension=f".{extension}",
            initialfile=filename,
            filetypes=[(f"{extension.upper()} files", f"*.{extension}")],
        )
        return path, include_total

    @staticmethod
    def _safe_filename(filename):
        return re.sub(r"[^A-Za-z0-9._-]+", "_", filename)

    def _print_data(self, include_total):
        if include_total or not self.has_total_score:
            return self.columns, self.records
        columns = self.columns[:-1]
        records = [row[:-1] for row in self.records]
        return columns, records

    def export_excel(self):
        path, include_total = self._ask_export_path("xlsx")
        if not path:
            return
        columns, records = self._print_data(include_total)
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
            from openpyxl.utils import get_column_letter

            workbook = Workbook()
            worksheet = workbook.active
            worksheet.title = self.report_title[:31]
            worksheet.append(
                [f"{self.report_title} | {self.grade} | {self.subject} | {self.term}"]
            )
            worksheet.merge_cells(
                start_row=1, start_column=1, end_row=1, end_column=len(columns)
            )
            worksheet.append(columns)
            for record in records:
                worksheet.append(record)
            worksheet.freeze_panes = "A3"
            worksheet.auto_filter.ref = (
                f"A2:{get_column_letter(len(columns))}{worksheet.max_row}"
            )
            worksheet.sheet_properties.pageSetUpPr.fitToPage = True
            worksheet.page_setup.orientation = "landscape"
            worksheet.page_setup.paperSize = worksheet.PAPERSIZE_A4
            worksheet.page_setup.fitToWidth = 1
            worksheet.page_setup.fitToHeight = 0
            worksheet.print_title_rows = "2:2"
            thin_border = Side(style="thin", color="B7C9C0")
            for cell in worksheet[1]:
                cell.font = Font(bold=True, color="FFFFFF", size=13)
                cell.fill = PatternFill("solid", fgColor="245B4B")
                cell.alignment = Alignment(horizontal="center")
            for cell in worksheet[2]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="245B4B")
                cell.alignment = Alignment(wrap_text=True, vertical="center")
            for row in worksheet.iter_rows(min_row=1):
                for cell in row:
                    cell.border = Border(bottom=thin_border)
                    if cell.row > 2:
                        cell.alignment = Alignment(vertical="center", wrap_text=True)
            for index in range(1, len(columns) + 1):
                worksheet.column_dimensions[get_column_letter(index)].width = (
                    25 if index == 1 else 30 if index == 2 else 13
                )
            workbook.save(path)
        except Exception as error:
            messagebox.showerror("Excel Export Failed", str(error), parent=self)
            return
        messagebox.showinfo("Export Complete", "Excel results saved.", parent=self)

    def export_pdf(self):
        path, include_total = self._ask_export_path("pdf")
        if not path:
            return
        columns, records = self._print_data(include_total)
        try:
            from reportlab.lib import colors
            from reportlab.lib.pagesizes import A4, landscape
            from reportlab.lib.units import mm
            from reportlab.platypus import Paragraph, SimpleDocTemplate, Table, TableStyle

            page_width = landscape(A4)[0]
            usable_width = page_width - 24 * mm
            weights = [1.5, 2.4] + [1.0] * max(0, len(columns) - 2)
            total_weight = sum(weights)
            column_widths = [usable_width * weight / total_weight for weight in weights]
            document = SimpleDocTemplate(
                path,
                pagesize=landscape(A4),
                rightMargin=12 * mm,
                leftMargin=12 * mm,
                topMargin=12 * mm,
                bottomMargin=12 * mm,
            )
            title = Paragraph(
                f"{self.report_title} | {self.grade} | {self.subject} | {self.term}",
                __import__("reportlab.lib.styles", fromlist=["getSampleStyleSheet"])
                .getSampleStyleSheet()["Title"],
            )
            table = Table([columns] + records, colWidths=column_widths, repeatRows=1)
            table.setStyle(
                TableStyle(
                    [
                        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#245B4B")),
                        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                        ("FONTSIZE", (0, 0), (-1, -1), 8),
                        ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#AEBBB5")),
                        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#EDF3EF")]),
                        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                    ]
                )
            )
            document.build([title, table])
        except Exception as error:
            messagebox.showerror("PDF Export Failed", str(error), parent=self)
            return
        messagebox.showinfo("Export Complete", "PDF results saved.", parent=self)